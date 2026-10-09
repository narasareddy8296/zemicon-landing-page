import io
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from flask import Flask
from openpyxl import Workbook, load_workbook

from app import app as application
from services.landing_price_v2 import landing_v2
from services.landing_price_v2.excel_import import inspect_upload


def make_master_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A4"] = "Manufacturer Part Number (MPN):"
    sheet["A5"] = "Component Category:"
    sheet["A6"] = "HSN / CTSH:"
    sheet["A28"] = "Basic Customs Duty (BCD) Rate %:"
    sheet["A30"] = "Social Welfare Surcharge (SWS) Rate %:"
    for column, record in enumerate(
        [
            ("ABC123", "Semiconductor", "85423100", 0.1, 0.1),
            ("PASSIVE1", "Passives", "85334090", 0, 0.1),
            ("STLINK-V3PWR", "Electromechanical", "84719000", 0, 0.1),
        ],
        2,
    ):
        for row, value in zip((4, 5, 6, 28, 30), record):
            sheet.cell(row=row, column=column, value=value)
    workbook.save(path)
    workbook.close()


class LandingPriceV2Tests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        master_path = root / "master.xlsx"
        make_master_workbook(master_path)
        self.app = Flask(
            __name__,
            template_folder=str(Path(__file__).resolve().parents[1] / "templates"),
            static_folder=str(Path(__file__).resolve().parents[1] / "static"),
            instance_path=str(root / "instance"),
        )
        self.app.config.update(
            TESTING=True,
            LANDING_V2_DATABASE=str(root / "landing.sqlite"),
            DIGIKEY_MASTER_FILE=str(master_path),
        )
        self.app.register_blueprint(landing_v2)
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def calculate(self, items, **overrides):
        payload = {
            "items": items,
            "remittance_applicable": "no",
            "cha_applicable": "no",
            "domestic_trucking_applicable": "no",
            "igst_enabled": "no",
            "margin_percent": 0,
            **overrides,
        }
        return self.client.post("/api/landing/v2/calculate", json=payload)

    def test_master_excel_seeds_lookup_with_percent_rates(self):
        response = self.client.get("/api/landing/v2/product/ abc123 ")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json(),
            {
                "found": True,
                "mpn": "ABC123",
                "category": "Semiconductor",
                "hsn": "85423100",
                "bcd": 10,
                "sws": 10,
                "source": "DigiKey Excel",
            },
        )

    def test_digikey_tier_selection_matches_frontend_break_and_base_fallback(self):
        from services.landing_price_v2.digikey_api import select_price_tier

        tiers = [
            {"break_quantity": 1, "unit_price": 57.33, "currency": "INR"},
            {"break_quantity": 10, "unit_price": 35.067, "currency": "INR"},
            {"break_quantity": 50, "unit_price": 26.2762, "currency": "INR"},
            {"break_quantity": 100, "unit_price": 23.5626, "currency": "INR"},
            {"break_quantity": 500, "unit_price": 18.99152, "currency": "INR"},
        ]
        for quantity, price_break in ((1, 1), (10, 10), (25, 10), (50, 50), (100, 100), (750, 500), (0, None)):
            selected = select_price_tier(tiers, quantity)
            self.assertEqual(selected["break_quantity"] if selected else None, price_break)
        self.assertEqual(select_price_tier(tiers, 0), None)
        self.assertIsNone(select_price_tier([{"break_quantity": 1, "unit_price": None, "currency": "INR"}], 1))
        self.assertIsNone(select_price_tier([], 1))

    def test_digikey_package_selection_uses_lowest_moq(self):
        from services.landing_price_v2.digikey_api import select_package_variation

        reel = {"package_type": "Tape & Reel", "minimum_order_quantity": 3000}
        cut_tape = {"package_type": "Cut Tape", "minimum_order_quantity": 1}
        self.assertIs(select_package_variation([reel, cut_tape]), cut_tape)
        self.assertIsNone(select_package_variation([]))

    def test_digikey_customer_tier_falls_back_to_standard_if_unusable(self):
        from services.landing_price_v2.digikey_api import _parse_variation

        variation = _parse_variation({
            "MyPricing": [{"BreakQuantity": 1, "UnitPrice": None}],
            "StandardPricing": [{"BreakQuantity": 1, "UnitPrice": 2.5}],
        }, "EUR")
        self.assertEqual(variation["pricing_source"], "StandardPricing")
        self.assertEqual(variation["pricing_tiers"][0]["unit_price"], 2.5)


    def test_catalog_hsn_is_returned_for_exact_mpn_match(self):
        response = self.client.get("/api/landing/v2/product/stlink-v3pwr")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["found"])
        self.assertEqual(response.get_json()["mpn"], "STLINK-V3PWR")
        self.assertEqual(response.get_json()["hsn"], "84719000")

    @patch("services.landing_price_v2.digikey_api.requests.request")
    def test_digikey_search_pricing_returns_observed_product_fields_and_currency(self, request_api):
        self.app.config.update(
            DIGIKEY_CLIENT_ID="test-client",
            DIGIKEY_CLIENT_SECRET="test-secret",
            DIGIKEY_CUSTOMER_ID="16159631",
        )
        token = unittest.mock.Mock(ok=True)
        token.json.return_value = {"access_token": "test-token", "expires_in": 600}
        pricing = unittest.mock.Mock(ok=True)
        pricing.json.return_value = {
            "ProductsCount": 1,
            "SettingsUsed": {"SearchLocale": {"Site": "IN", "Language": "en", "Currency": "EUR"}},
            "ProductPricings": [{
                "ManufacturerProductNumber": "ABC123",
                "Manufacturer": {"Name": "Example Manufacturer"},
                "DigiKeyProductNumber": "DIGIKEY-SKU-ND",
                "Category": {"Name": "Integrated Circuits"},
                "Description": {"ProductDescription": "Example component", "DetailedDescription": "A detailed example"},
                "QuantityAvailable": 99,
                "ProductVariations": [{
                    "DigiKeyProductNumber": "DIGIKEY-SKU-ND",
                    "PackageType": {"Name": "Cut Tape (CT) & Digi-Reel®"},
                    "QuantityAvailableforPackageType": 12,
                    "MinimumOrderQuantity": 1,
                    "StandardPricing": [
                        {"BreakQuantity": 1, "UnitPrice": 5.2, "TotalPrice": 5.2},
                        {"BreakQuantity": 10, "UnitPrice": 4.0, "TotalPrice": 40.0},
                        {"BreakQuantity": 25, "UnitPrice": 3.5, "TotalPrice": 87.5},
                    ],
                    "MyPricing": [],
                }],
            }],
        }
        request_api.side_effect = [token, pricing]

        response = self.client.get("/api/landing/v2/digikey-product/ABC123?currency=USD")
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertTrue(result["found"])
        self.assertEqual(result["currency"], "EUR")
        match = result["matches"][0]
        self.assertEqual(match["manufacturer"], "Example Manufacturer")
        self.assertEqual(match["category"], "Integrated Circuits")
        self.assertEqual(match["description"], "Example component")
        self.assertEqual(match["detailed_description"], "A detailed example")
        self.assertEqual(match["available_quantity"], 99)
        self.assertEqual(match["variations"][0]["available_quantity"], 12)
        self.assertEqual(match["variations"][0]["package_type"], "Cut Tape (CT) & Digi-Reel®")
        self.assertEqual([tier["break_quantity"] for tier in match["variations"][0]["pricing_tiers"]], [1, 10, 25])
        self.assertEqual(request_api.call_count, 2)
        self.assertTrue(request_api.call_args_list[1].args[1].endswith("/search/ABC123/pricing"))
        self.assertEqual(request_api.call_args_list[1].kwargs["headers"]["X-DIGIKEY-Locale-Currency"], "USD")
        self.assertEqual(request_api.call_args_list[1].kwargs["headers"]["X-DIGIKEY-Customer-Id"], "16159631")
        self.assertNotIn("test-secret", response.get_data(as_text=True))

    @patch("services.landing_price_v2.digikey_api.requests.request")
    def test_digikey_customer_pricing_tiers_are_used_when_present(self, request_api):
        self.app.config.update(DIGIKEY_CLIENT_ID="test-client", DIGIKEY_CLIENT_SECRET="test-secret")
        token = unittest.mock.Mock(ok=True)
        token.json.return_value = {"access_token": "test-token", "expires_in": 600}
        pricing = unittest.mock.Mock(ok=True)
        pricing.json.return_value = {
            "SettingsUsed": {"SearchLocale": {"Currency": "INR"}},
            "ProductPricings": [{
                "ManufacturerProductNumber": "ABC123",
                "ProductVariations": [{
                    "MinimumOrderQuantity": 1,
                    "MyPricing": [{"BreakQuantity": 1, "UnitPrice": 12.5}],
                    "StandardPricing": [{"BreakQuantity": 1, "UnitPrice": 15.0}],
                }],
            }],
        }
        request_api.side_effect = [token, pricing]
        result = self.client.get("/api/landing/v2/digikey-product/ABC123").get_json()
        variation = result["matches"][0]["variations"][0]
        self.assertEqual(variation["pricing_source"], "MyPricing")
        self.assertEqual(variation["pricing_tiers"][0]["unit_price"], 12.5)

    def test_live_digikey_endpoint_reports_not_found_without_fabricated_fields(self):
        from services.landing_price_v2.digikey_api import lookup_product_details

        with self.app.app_context():
            self.app.config.update(DIGIKEY_CLIENT_ID="test-client", DIGIKEY_CLIENT_SECRET="test-secret")
            with patch("services.landing_price_v2.digikey_api.requests.request") as request_api:
                token = unittest.mock.Mock(ok=True)
                token.json.return_value = {"access_token": "test-token", "expires_in": 600}
                pricing = unittest.mock.Mock(ok=True)
                pricing.json.return_value = {"ProductsCount": 0, "ProductPricings": [], "SettingsUsed": {"SearchLocale": {"Site": "IN", "Language": "en", "Currency": "INR"}}}
                request_api.side_effect = [token, pricing]
                result = lookup_product_details("UNKNOWN")
        self.assertEqual(result["matches"], [])
        self.assertEqual(result["currency"], "INR")

    @patch("services.landing_price_v2.digikey_api.requests.request")
    def test_digikey_authentication_rate_limit_timeout_and_currency_failures_are_reported(self, request_api):
        self.app.config.update(DIGIKEY_CLIENT_ID="test-client", DIGIKEY_CLIENT_SECRET="test-secret")
        token = unittest.mock.Mock(ok=True)
        token.json.return_value = {"access_token": "test-token", "expires_in": 600}
        unauthorized = unittest.mock.Mock(ok=False, status_code=401)
        request_api.side_effect = [token, unauthorized]
        response = self.client.get("/api/landing/v2/digikey-product/ABC123")
        self.assertEqual(response.status_code, 502)
        self.assertIn("authentication", response.get_json()["error"].lower())

        self.app.extensions.pop("digikey_oauth_token", None)
        request_api.side_effect = [token, unittest.mock.Mock(ok=False, status_code=429)]
        response = self.client.get("/api/landing/v2/digikey-product/ABC123")
        self.assertEqual(response.status_code, 429)
        self.assertIn("rate limit", response.get_json()["error"].lower())

        from requests import Timeout
        self.app.extensions.pop("digikey_oauth_token", None)
        request_api.side_effect = [token, Timeout("timeout")]
        response = self.client.get("/api/landing/v2/digikey-product/ABC123")
        self.assertEqual(response.status_code, 504)

    def test_freight_threshold_and_item_tariffs(self):
        below = self.calculate(
            [{"mpn": "ABC123", "quantity": 10, "unit_price": 600, "currency": "INR"}]
        ).get_json()
        self.assertEqual(below["totals"]["international_freight"], 1200)
        self.assertEqual(below["items"][0]["insurance"], 67.5)
        self.assertEqual(below["items"][0]["assessable_value"], 7267.5)
        self.assertEqual(below["items"][0]["bcd_amount"], 727)
        self.assertEqual(below["items"][0]["sws_amount"], 72.7)
        self.assertEqual(below["items"][0]["base_landed_cost"], 8067.2)

        threshold = self.calculate(
            [{"mpn": "ABC123", "quantity": 10, "unit_price": 700, "currency": "INR"}]
        ).get_json()
        self.assertEqual(threshold["totals"]["international_freight"], 0)

    def test_freight_uses_shipment_threshold_and_invoice_share_allocation(self):
        result = self.calculate(
            [
                {"mpn": "ABC123", "quantity": 1, "unit_price": 8000, "currency": "INR"},
                {"mpn": "PASSIVE1", "quantity": 1, "unit_price": 2000, "currency": "INR"},
            ]
        ).get_json()
        self.assertEqual(result["totals"]["international_freight"], 0)

        below_threshold = self.calculate(
            [
                {"mpn": "ABC123", "quantity": 1, "unit_price": 4000, "currency": "INR"},
                {"mpn": "PASSIVE1", "quantity": 1, "unit_price": 2000, "currency": "INR"},
            ]
        ).get_json()
        self.assertEqual(below_threshold["items"][0]["international_freight"], 800)
        self.assertEqual(below_threshold["items"][1]["international_freight"], 400)
        self.assertEqual(below_threshold["totals"]["international_freight"], 1200)

    def test_shipment_charges_are_allocated_once_across_lines(self):
        response = self.calculate(
            [
                {"mpn": "ABC123", "quantity": 1, "unit_price": 6000, "currency": "INR"},
                {"mpn": "PASSIVE1", "quantity": 1, "unit_price": 4000, "currency": "INR"},
            ],
            remittance_applicable="yes",
            cha_applicable="yes",
            domestic_trucking_applicable="yes",
        )
        result = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(result["totals"]["total_shipment_charges"], 4712.5)
        self.assertEqual(result["items"][0]["international_freight"], 0)
        self.assertEqual(result["items"][1]["international_freight"], 0)
        self.assertEqual(result["items"][0]["remittance"], 1800)
        self.assertEqual(result["items"][1]["remittance"], 1200)
        self.assertEqual(
            sum(item["cha_port_dues"] for item in result["items"]),
            1100,
        )
        self.assertEqual(
            sum(item["domestic_trucking"] for item in result["items"]),
            500,
        )

    def test_currency_conversion_applies_before_threshold(self):
        response = self.calculate(
            [{
                "mpn": "PASSIVE1",
                "quantity": 1,
                "unit_price": 70,
                "currency": "USD",
                "live_exchange_rate": 100,
            }]
        )
        result = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(result["items"][0]["live_exchange_rate_inr"], 100)
        self.assertEqual(result["items"][0]["exchange_rate_inr"], 102)
        self.assertEqual(result["items"][0]["line_invoice_currency"], 70)
        self.assertEqual(result["items"][0]["line_invoice_inr"], 7140)
        self.assertEqual(result["totals"]["total_invoice_inr"], 7140)
        self.assertEqual(result["totals"]["international_freight"], 0)

    def test_usd_live_rate_gets_two_percent_and_inr_does_not(self):
        usd = self.calculate(
            [{
                "mpn": "PASSIVE1",
                "quantity": 2,
                "unit_price": 10,
                "currency": "USD",
                "live_exchange_rate": 80,
            }]
        ).get_json()["items"][0]
        self.assertEqual(usd["exchange_rate_inr"], 81.6)
        self.assertEqual(usd["unit_price_inr"], 816)
        self.assertEqual(usd["line_invoice_currency"], 20)
        self.assertEqual(usd["line_invoice_inr"], 1632)

        inr = self.calculate(
            [{"mpn": "PASSIVE1", "quantity": 2, "unit_price": 10, "currency": "INR"}]
        ).get_json()["items"][0]
        self.assertIsNone(inr["live_exchange_rate_inr"])
        self.assertEqual(inr["exchange_rate_inr"], 1)
        self.assertEqual(inr["unit_price_inr"], 10)
        self.assertEqual(inr["line_invoice_currency"], 20)
        self.assertEqual(inr["line_invoice_inr"], 20)

    def test_usd_requires_live_interbank_rate(self):
        response = self.calculate(
            [{"mpn": "PASSIVE1", "quantity": 1, "unit_price": 10, "currency": "USD"}]
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("live interbank USD/INR", response.get_json()["error"])

    def test_igst_is_optional_and_applied_after_margin(self):
        item = [{"mpn": "PASSIVE1", "quantity": 1, "unit_price": 7000, "currency": "INR"}]
        off = self.calculate(item).get_json()
        self.assertFalse(off["totals"]["igst_enabled"])
        self.assertNotIn("igst", off["items"][0])
        self.assertEqual(off["items"][0]["margin_amount"], 0)
        self.assertEqual(off["items"][0]["margin_inclusive_value"], 7078.75)

        on = self.calculate(item, igst_enabled="yes", margin_percent=10).get_json()
        self.assertEqual(on["items"][0]["base_landed_cost"], 7078.75)
        self.assertEqual(on["items"][0]["margin_amount"], 707.88)
        self.assertEqual(on["items"][0]["margin_inclusive_value"], 7786.63)
        self.assertEqual(on["items"][0]["igst"], 1401.59)
        self.assertEqual(on["items"][0]["final_price_including_igst"], 9188.22)

    def test_margin_applies_to_landed_cost_and_each_line_without_igst(self):
        result = self.calculate(
            [
                {"mpn": "ABC123", "quantity": 2, "unit_price": 1000, "currency": "INR"},
                {"mpn": "PASSIVE1", "quantity": 1, "unit_price": 2000, "currency": "INR"},
            ],
            margin_percent=10,
        ).get_json()

        first, second = result["items"]
        self.assertFalse(result["totals"]["igst_enabled"])
        self.assertEqual(first["base_landed_cost"], 2910.7)
        self.assertEqual(first["margin_amount"], 291.07)
        self.assertEqual(first["margin_inclusive_value"], 3201.77)
        self.assertEqual(first["per_unit_selling_price"], 1600.89)
        self.assertEqual(second["base_landed_cost"], 2622.5)
        self.assertEqual(second["margin_amount"], 262.25)
        self.assertEqual(second["margin_inclusive_value"], 2884.75)
        self.assertEqual(result["totals"]["base_landed_cost"], 5533.2)
        self.assertEqual(result["totals"]["margin_amount"], 553.32)
        self.assertEqual(result["totals"]["margin_inclusive_value"], 6086.52)

    def test_missing_master_record_can_be_added_and_reused(self):
        request_data = {
            "items": [{"mpn": "NEW-1", "quantity": 1, "unit_price": 100, "currency": "INR"}]
        }
        missing = self.client.post("/api/landing/v2/calculate", json=request_data)
        self.assertEqual(missing.status_code, 400)
        self.assertIn("BCD rate", missing.get_json()["error"])

        added = self.client.post(
            "/api/landing/v2/master",
            json={
                "mpn": "NEW-1",
                "category": "Semiconductor",
                "hsn": "85423100",
                "bcd": "0%",
                "sws": "10%",
            },
        )
        self.assertEqual(added.status_code, 200)
        self.assertEqual(added.get_json()["created"], ["NEW-1"])
        repeated = self.client.get("/api/landing/v2/product/new-1")
        self.assertTrue(repeated.get_json()["found"])
        self.assertEqual(self.client.post("/api/landing/v2/calculate", json=request_data).status_code, 200)

    def test_unknown_mpn_can_calculate_with_entered_bcd_and_sws(self):
        response = self.calculate(
            [{
                "mpn": "CUSTOM-1",
                "quantity": 1,
                "unit_price": 1000,
                "currency": "INR",
                "category": "Custom category",
                "bcd_rate": 5,
                "sws_rate": 10,
            }]
        )
        result = response.get_json()
        self.assertEqual(response.status_code, 200)
        item = result["items"][0]
        self.assertEqual(item["category"], "Custom category")
        self.assertEqual(item["insurance"], 11.25)
        self.assertEqual(item["assessable_value"], 2211.25)
        self.assertEqual(item["bcd_amount"], 111)
        self.assertEqual(item["sws_amount"], 11.1)
        self.assertEqual(item["base_landed_cost"], 2333.35)
        self.assertFalse(self.client.get("/api/landing/v2/product/CUSTOM-1").get_json()["found"])

    def test_other_charges_allocated_by_invoice_and_excluded_from_assessable_value(self):
        result = self.calculate(
            [
                {"mpn": "ABC123", "quantity": 1, "unit_price": 6000, "currency": "INR"},
                {"mpn": "PASSIVE1", "quantity": 1, "unit_price": 4000, "currency": "INR"},
            ],
            other_charges_total=1000,
        ).get_json()
        first, second = result["items"]
        self.assertEqual(first["other_charges"], 600)
        self.assertEqual(second["other_charges"], 400)
        self.assertEqual(first["insurance"], 67.5)
        self.assertEqual(second["insurance"], 45)
        self.assertEqual(first["international_freight"], 0)
        self.assertEqual(second["international_freight"], 0)
        self.assertEqual(first["assessable_value"], 6067.5)
        self.assertEqual(second["assessable_value"], 4045)
        self.assertEqual(result["totals"]["insurance"], 112.5)
        self.assertEqual(result["totals"]["other_charges"], 1000)
        self.assertEqual(result["totals"]["total_shipment_charges"], 1112.5)
        self.assertEqual(result["totals"]["assessable_value"], 10112.5)

    def test_master_addition_does_not_overwrite_existing_tariff(self):
        response = self.client.post(
            "/api/landing/v2/master",
            json={
                "mpn": "ABC123",
                "category": "Changed category",
                "hsn": "85423100",
                "bcd": 25,
                "sws": 0,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["already_present"], ["ABC123"])
        product = self.client.get("/api/landing/v2/product/ABC123").get_json()
        self.assertEqual(product["category"], "Semiconductor")
        self.assertEqual(product["bcd"], 10)

    def test_excel_import_endpoint_ignores_extended_price(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Description", "MPN", "Qty", "Unit Price", "Total Price", "Notes"])
        sheet.append(["Part", "ABC123", 4, 2.5, 10, "ignored"])
        contents = io.BytesIO()
        workbook.save(contents)
        workbook.close()
        contents.seek(0)
        response = self.client.post(
            "/api/landing/v2/import-excel",
            data={"file": (contents, "purchase.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        item = response.get_json()["items"][0]
        self.assertEqual(item["mpn"], "ABC123")
        self.assertEqual(item["unit_price"], 2.5)
        self.assertEqual(item["quantity"], 4)
        self.assertEqual(item["currency"], "USD")
        self.assertTrue(item["master"]["found"])

    @patch("services.landing_price_v2.routes.lookup_product_details")
    def test_excel_import_response_contains_live_details_and_quantity_price(self, lookup):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Unit Price", "Quantity"])
        sheet.append(["ABC123", 2.5, 4])
        contents = io.BytesIO()
        workbook.save(contents)
        workbook.close()
        contents.seek(0)
        lookup.return_value = {
            "currency": "USD",
            "matches": [{
                "manufacturer_part_number": "ABC123",
                "manufacturer": "Example Manufacturer",
                "description": "Example component",
                "category": "Integrated Circuits",
                "digikey_part_number": "DK-ABC123",
                "exact_mpn_match": True,
                "variations": [{
                    "digikey_product_number": "DK-ABC123-CT",
                    "package_type": "Cut Tape",
                    "minimum_order_quantity": 1,
                    "available_quantity": 40,
                    "pricing_tiers": [
                        {"break_quantity": 1, "unit_price": 2.5, "currency": "USD"},
                        {"break_quantity": 4, "unit_price": 2.0, "currency": "USD"},
                    ],
                }],
            }],
        }

        response = self.client.post(
            "/api/landing/v2/import-excel",
            data={"file": (contents, "purchase.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 200)
        item = response.get_json()["items"][0]
        self.assertTrue(item["master"]["found"])
        self.assertEqual(item["mpn"], "ABC123")
        self.assertEqual(item["product_details"]["matches"][0]["manufacturer"], "Example Manufacturer")
        self.assertEqual(item["live_quote"]["unit_price"], 2.0)
        self.assertEqual(item["live_quote"]["requested_quantity"], 4)
        self.assertEqual(item["unit_price"], 2.5)  # BOM value is preserved.
        lookup.assert_called_once_with("ABC123", requested_currency="USD")

    def test_excel_import_rejects_currencies_other_than_usd_or_inr(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Unit Price", "Currency"])
        sheet.append(["ABC123", 1.5, "EUR"])
        contents = io.BytesIO()
        workbook.save(contents)
        workbook.close()
        contents.seek(0)
        response = self.client.post(
            "/api/landing/v2/import-excel",
            data={"file": (contents, "purchase.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("USD or INR", response.get_json()["error"])

    def test_excel_import_allows_manual_header_and_column_selection(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Part Reference", "Per Item Cost", "Quantity"])
        sheet.append(["ABC123", 2.5, 4])
        contents = io.BytesIO()
        workbook.save(contents)
        workbook.close()

        def upload(fields=None):
            return self.client.post(
                "/api/landing/v2/import-excel",
                data={
                    **(fields or {}),
                    "default_currency": "USD",
                    "file": (io.BytesIO(contents.getvalue()), "purchase.xlsx"),
                },
                content_type="multipart/form-data",
            )

        workbook_choice = upload()
        self.assertTrue(workbook_choice.get_json()["needs_selection"])
        selected_header = workbook_choice.get_json()["sheets"][0]
        column_choice = upload(selected_header)
        self.assertTrue(column_choice.get_json()["needs_selection"])
        self.assertEqual(
            [header["label"] for header in column_choice.get_json()["headers"]],
            ["Part Reference", "Per Item Cost", "Quantity"],
        )
        imported = upload(
            {
                **selected_header,
                "mpn_column": 0,
                "unit_price_column": 1,
                "quantity_column": 2,
            }
        )
        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.get_json()["items"][0]["unit_price"], 2.5)
        self.assertEqual(imported.get_json()["items"][0]["quantity"], 4)
        self.assertEqual(imported.get_json()["items"][0]["currency"], "USD")

    def test_v2_screen_and_master_export(self):
        page = self.client.get("/landing/v2")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Zemicon | Landing Price Calculator", page.data)
        self.assertIn(b'<option value="digikey" selected>DigiKey</option>', page.data)
        self.assertIn(b'<option value="other">Other</option>', page.data)
        self.assertIn(b'id="invoiceCurrency"', page.data)
        self.assertIn(b'id="liveRate"', page.data)
        self.assertNotIn(b'<th>Currency</th>', page.data)
        self.assertNotIn(b'<th>Live interbank USD/INR rate</th>', page.data)
        self.assertIn(b'<th>Amount (INR \xe2\x82\xb9)</th>', page.data)
        self.assertIn(b'id="invoiceSummary"', page.data)
        self.assertIn(b'id="tariffsSection"', page.data)
        self.assertIn(b'id="tariffRows"', page.data)
        self.assertLess(
            page.data.index(b'id="invoiceSummary"'),
            page.data.index(b'id="tariffsSection"'),
        )
        exported = self.client.get("/api/landing/v2/master/export")
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(
            exported.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        exported_workbook = load_workbook(io.BytesIO(exported.data), read_only=True)
        self.assertEqual(
            list(next(exported_workbook.active.iter_rows(values_only=True))),
            ["MPN", "Category", "HSN / CTSH", "BCD (%)", "SWS (%)", "Source"],
        )
        exported_workbook.close()

    def test_root_redirects_to_v2_and_legacy_endpoints_are_removed(self):
        client = application.test_client()
        root = client.get("/")
        self.assertEqual(root.status_code, 302)
        self.assertEqual(root.headers["Location"], "/landing/v2")
        self.assertEqual(client.post("/api/calculate", json={}).status_code, 404)
        self.assertEqual(client.get("/api/dhl/rate").status_code, 404)

    def test_ambiguous_excel_columns_are_not_guessed(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["MPN", "Manufacturer Part Number", "Unit Price", "Price Per Unit", "Qty"])
        sheet.append(["A", "A", 2.5, 2.5, 1])
        contents = io.BytesIO()
        workbook.save(contents)
        workbook.close()
        contents.seek(0)
        result = inspect_upload(contents)
        self.assertTrue(result["needs_selection"])
        self.assertIn("unit_price", result["candidates"])

        contents.seek(0)
        imported = inspect_upload(
            contents,
            {
                "sheet": "Sheet",
                "header_row": 1,
                "mpn_column": 0,
                "unit_price_column": 2,
                "quantity_column": 4,
                "currency_column": "",
            },
        )
        self.assertEqual(imported["items"][0]["mpn"], "A")
        self.assertEqual(imported["items"][0]["unit_price"], 2.5)


if __name__ == "__main__":
    unittest.main()
