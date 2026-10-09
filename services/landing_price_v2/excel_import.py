import csv
import re
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


MAX_IMPORT_ROWS = 5000
MAX_HEADER_SCAN_ROWS = 30
MPN_ALIASES = {
    "mpn",
    "manufacturerpartnumber",
    "manufacturerpartno",
    "manufacturerpart",
    "manufacturerpn",
    "mfrpart",
    "mfrpartnumber",
    "mfrpartnum",
    "mfrpartno",
    "mfrpn",
    "partnumber",
    "partno",
    "part",
}
PRICE_ALIASES = {
    "unitprice",
    "unitpriceusd",
    "unitpriceinr",
    "priceeach",
    "priceusd",
    "priceinr",
    "eachprice",
    "perunitprice",
    "unitcost",
    "unitcostusd",
    "unitcostinr",
    "costeach",
    "priceunit",
    "priceperunit",
    "purchaseprice",
    "vendorunitprice",
    "rate",
    "unitrate",
    "rateeach",
    "price",
    "cost",
}
QUANTITY_ALIASES = {
    "quantity",
    "qty",
    "bomqty",
    "orderquantity",
    "orderqty",
    "reqqty",
    "requiredqty",
    "purchaseqty",
}
CURRENCY_ALIASES = {"currency", "billingcurrency", "pricecurrency"}
SUPPORTED_UPLOAD_FORMATS = {".xlsx", ".xls", ".csv", ".tsv", ".txt", ".pdf"}


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _rows_to_xlsx(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "BOM"
    for row in rows:
        worksheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def normalize_upload(stream, filename):
    """Return an xlsx stream for supported spreadsheet, CSV, and text PDF files."""
    suffix = Path(filename).suffix.casefold()
    if suffix not in SUPPORTED_UPLOAD_FORMATS:
        raise ValueError("Upload a .xlsx, .xls, .csv, .tsv, .txt, or text-based .pdf BOM file.")
    if suffix == ".xlsx":
        return stream

    content = stream.read()
    if suffix in {".csv", ".tsv", ".txt"}:
        decoded = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("Could not decode this CSV file. Save it as UTF-8 and upload again.")
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        return _rows_to_xlsx(csv.reader(StringIO(decoded), dialect))

    if suffix == ".xls":
        import xlrd

        try:
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            rows = (
                book.sheet_by_index(0).row_values(index)
                for index in range(book.sheet_by_index(0).nrows)
            )
            output = _rows_to_xlsx(rows)
            book.release_resources()
            return output
        except Exception as exc:
            raise ValueError("Could not read this .xls workbook. Save it as .xlsx and try again.") from exc

    try:
        import pdfplumber

        extracted_rows = []
        with pdfplumber.open(BytesIO(content)) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables() or []:
                    extracted_rows.extend(table)
            if not extracted_rows:
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    for line in text.splitlines():
                        cells = re.split(r"\s{2,}|\s*\|\s*", line.strip())
                        if len(cells) > 1:
                            extracted_rows.append(cells)
        if not extracted_rows:
            raise ValueError(
                "No readable table was found in this PDF. Scanned PDFs need OCR; export the BOM as CSV or Excel and upload it."
            )
        return _rows_to_xlsx(extracted_rows)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this PDF. Export the BOM as CSV or Excel and upload it.") from exc


def _header_candidates(values):
    candidates = {"mpn": [], "unit_price": [], "quantity": [], "currency": []}
    labels = []
    for column, value in enumerate(values, 1):
        if value is None or not str(value).strip():
            continue
        label = " ".join(str(value or "").split())
        labels.append({"index": column - 1, "label": label or f"Column {get_column_letter(column)}"})
        normalized = normalize_header(value)
        is_extended_price = any(term in normalized for term in ("extended", "total", "amount"))
        if normalized in MPN_ALIASES:
            candidates["mpn"].append(column - 1)
        if normalized in PRICE_ALIASES and not is_extended_price:
            candidates["unit_price"].append(column - 1)
        if normalized in QUANTITY_ALIASES:
            candidates["quantity"].append(column - 1)
        if normalized in CURRENCY_ALIASES:
            candidates["currency"].append(column - 1)
    return candidates, labels


def _scan_workbook(workbook):
    detected = []
    for worksheet in workbook.worksheets:
        for row_number in range(1, min(worksheet.max_row, MAX_HEADER_SCAN_ROWS) + 1):
            values = next(
                worksheet.iter_rows(
                    min_row=row_number,
                    max_row=row_number,
                    values_only=True,
                )
            )
            candidates, labels = _header_candidates(values)
            if (
                candidates["mpn"]
                or candidates["unit_price"]
                or sum(
                    1 for value in values
                    if value is not None and str(value).strip()
                ) >= 2
            ):
                detected.append(
                    {
                        "sheet": worksheet.title,
                        "header_row": row_number,
                        "candidates": candidates,
                        "headers": labels,
                    }
                )
    return detected


def inspect_upload(stream, selections=None):
    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        default_currency = (selections or {}).get("default_currency", "USD")
        selection_keys = (
            "sheet",
            "header_row",
            "mpn_column",
            "unit_price_column",
            "quantity_column",
            "currency_column",
        )
        if not any(
            (selections or {}).get(key) not in (None, "")
            for key in selection_keys
        ):
            selections = None
        detected = _scan_workbook(workbook)
        if selections is None or not (
            selections.get("mpn_column") not in (None, "")
            and selections.get("unit_price_column") not in (None, "")
        ):
            relevant = [
                entry
                for entry in detected
                if entry["candidates"]["mpn"] and entry["candidates"]["unit_price"]
            ]
            if selections is None and len(relevant) != 1:
                return {
                    "needs_selection": True,
                    "sheets": [
                        {"sheet": entry["sheet"], "header_row": entry["header_row"]}
                        for entry in relevant or detected
                    ],
                    "message": (
                        "Select a worksheet and header row."
                        if len(relevant) > 1
                        else "Could not uniquely identify a worksheet and header row."
                    ),
                }
            if selections is not None:
                selected = next(
                    (
                        entry for entry in detected
                        if entry["sheet"] == selections.get("sheet")
                        and entry["header_row"] == int(selections.get("header_row", 0))
                    ),
                    None,
                )
                if selected is None:
                    raise ValueError("The selected worksheet or header row is invalid.")
            else:
                selected = relevant[0]
            candidates = selected["candidates"]
            if (
                len(candidates["mpn"]) != 1
                or len(candidates["unit_price"]) != 1
                or len(candidates["quantity"]) > 1
                or len(candidates["currency"]) > 1
            ) and (
                selections is None
                or selections.get("mpn_column") in (None, "")
                or selections.get("unit_price_column") in (None, "")
            ):
                return {
                    "needs_selection": True,
                    "sheets": [{"sheet": selected["sheet"], "header_row": selected["header_row"]}],
                    "headers": selected["headers"],
                    "candidates": candidates,
                    "message": "Select the correct columns before importing.",
                }
            selections = dict(selections or {})
            selections["default_currency"] = default_currency
            selections.update(
                {
                    "sheet": selected["sheet"],
                    "header_row": selected["header_row"],
                    "mpn_column": selections.get("mpn_column", candidates["mpn"][0]),
                    "unit_price_column": selections.get(
                        "unit_price_column", candidates["unit_price"][0]
                    ),
                    "quantity_column": selections.get(
                        "quantity_column",
                        candidates["quantity"][0] if candidates["quantity"] else None,
                    ),
                    "currency_column": selections.get(
                        "currency_column",
                        candidates["currency"][0] if candidates["currency"] else None,
                    ),
                }
            )

        worksheet = workbook[selections["sheet"]]
        header_row = int(selections["header_row"])
        headers = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            )
        )
        selected_columns = {
            key: (None if selections.get(key) in (None, "") else int(selections[key]))
            for key in ("mpn_column", "unit_price_column", "quantity_column", "currency_column")
        }
        if selected_columns["mpn_column"] is None or selected_columns["unit_price_column"] is None:
            raise ValueError("Select both an MPN column and a Per Unit Price column.")
        if not (0 <= selected_columns["mpn_column"] < len(headers)) or not (
            0 <= selected_columns["unit_price_column"] < len(headers)
        ):
            raise ValueError("The selected MPN or Per Unit Price column is invalid.")

        items = []
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            mpn = row[selected_columns["mpn_column"]] if selected_columns["mpn_column"] < len(row) else None
            if mpn is None or not str(mpn).strip():
                continue
            if len(items) >= MAX_IMPORT_ROWS:
                raise ValueError(f"Excel files are limited to {MAX_IMPORT_ROWS} line items.")
            price = row[selected_columns["unit_price_column"]] if selected_columns["unit_price_column"] < len(row) else None
            quantity = (
                row[selected_columns["quantity_column"]]
                if selected_columns["quantity_column"] is not None
                and selected_columns["quantity_column"] < len(row)
                else 1
            )
            currency = (
                row[selected_columns["currency_column"]]
                if selected_columns["currency_column"] is not None
                and selected_columns["currency_column"] < len(row)
                else selections.get("default_currency", "USD")
            )
            items.append(
                {
                    "mpn": " ".join(str(mpn).split()),
                    "unit_price": price,
                    "quantity": quantity if quantity is not None else 1,
                    "currency": str(currency or selections.get("default_currency", "USD")).strip().upper(),
                }
            )
        if not items:
            raise ValueError("No line items with MPN values were found below the selected header row.")

        return {
            "needs_selection": False,
            "sheet": worksheet.title,
            "header_row": header_row,
            "items": items,
        }
    finally:
        workbook.close()


def read_seed_master(path):
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        field_rows = {}
        label_fields = {
            "mpn": ("manufacturer part number",),
            "category": ("component category",),
            "hsn": ("hsn", "ctsh"),
            "bcd": ("basic customs duty (bcd) rate",),
            "sws": ("social welfare surcharge (sws) rate",),
        }
        for row_number in range(1, worksheet.max_row + 1):
            label = worksheet.cell(row=row_number, column=1).value
            normalized = normalize_header(label)
            for field, phrases in label_fields.items():
                if any(normalize_header(phrase) in normalized for phrase in phrases):
                    field_rows[field] = row_number
        if set(field_rows) != set(label_fields):
            raise ValueError("DigiKey master workbook is missing MPN, category, HSN, BCD, or SWS fields.")
        rows = []
        row_values = {}
        for field, row_number in field_rows.items():
            row_values[field] = next(
                worksheet.iter_rows(
                    min_row=row_number,
                    max_row=row_number,
                    min_col=2,
                    values_only=True,
                )
            )
        rows = []
        for offset in range(max(map(len, row_values.values()))):
            record = {
                field: values[offset] if offset < len(values) else None
                for field, values in row_values.items()
            }
            if not record["mpn"]:
                continue
            if not all(record.get(field) not in (None, "") for field in ("category", "hsn", "bcd", "sws")):
                continue
            rows.append(record)
        if not rows:
            raise ValueError("No complete DigiKey product records were found in the master workbook.")
        return rows
    finally:
        workbook.close()
